"""Bounded local ingestion. Inspection never launches Origin or sends the full dataset."""

import csv
import hashlib
import io
import math
import os
import tempfile
import uuid
import zipfile
from pathlib import Path

from .storage import Store, json_bytes, read_json, sha256, write_json

MAX_BYTES = 25 * 1024 * 1024
MAX_ROWS = 250_000
MAX_COLS = 128
MAX_CELLS = 2_000_000


def _rows(path: Path, sheet_name: str | None):
    if path.suffix == ".xlsx":
        import openpyxl

        with zipfile.ZipFile(path) as archive:
            infos = archive.infolist()
            if len(infos) > 2000 or sum(i.file_size for i in infos) > 100 * 1024 * 1024:
                raise ValueError("XLSX expanded size exceeds limit")
        book = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            if sheet_name is None and len(book.sheetnames) != 1:
                raise ValueError(f"Choose sheet_name from: {book.sheetnames[:20]}")
            if sheet_name is not None and sheet_name not in book.sheetnames:
                raise ValueError(
                    f"Worksheet {sheet_name!r} not found. Choose sheet_name from: {book.sheetnames[:20]}"
                )
            sheet = book[sheet_name if sheet_name is not None else book.sheetnames[0]]
            for row in sheet.iter_rows():
                if any(c.data_type == "f" for c in row):
                    raise ValueError(
                        "XLSX formulas need an explicit values-only export; no stale cached values used"
                    )
                yield ["" if c.value is None else str(c.value) for c in row]
        finally:
            book.close()
    else:
        if sheet_name:
            raise ValueError("sheet_name is only applicable to XLSX")
        raw = path.read_bytes()
        encoding = "utf-16" if raw.startswith((b"\xff\xfe", b"\xfe\xff")) else "utf-8-sig"
        content = raw.decode(encoding, errors="strict")
        delimiter = "\t" if path.suffix == ".tsv" else ","
        yield from csv.reader(io.StringIO(content, newline=""), delimiter=delimiter, strict=True)


def load_table(path: Path, sheet_name: str | None = None):
    iterator = iter(_rows(path, sheet_name))
    headers = next(iterator, [])
    if not 2 <= len(headers) <= MAX_COLS:
        raise ValueError(f"Expected 2 to {MAX_COLS} columns with a header row")
    if any(not h or len(h) > 160 or h != h.strip() for h in headers) or len(set(headers)) != len(headers):
        raise ValueError(
            "Column headers must be unique, nonempty, <=160 characters, without surrounding spaces"
        )
    rows = []
    for index, row in enumerate(iterator, 1):
        if index > MAX_ROWS or index * len(headers) > MAX_CELLS:
            raise ValueError("Data exceeds row/cell limit")
        if len(row) != len(headers):
            raise ValueError(f"Row {index + 1} has {len(row)} cells; expected {len(headers)}")
        if any(len(value) > 4096 for value in row):
            raise ValueError(f"Row {index + 1} contains a cell longer than 4096 characters")
        rows.append(row)
    if not rows:
        raise ValueError("Dataset has no data rows")
    return headers, rows


def inspect_dataset(store: Store, path: str, sheet_name: str | None = None) -> dict:
    source = store.input_path(path)
    fd, temporary = tempfile.mkstemp(suffix=source.suffix.lower(), dir=store.root / "datasets")
    temp = Path(temporary)
    try:
        with source.open("rb") as stream, os.fdopen(fd, "wb") as target:
            size = 0
            digest = hashlib.sha256()
            while chunk := stream.read(1024 * 1024):
                size += len(chunk)
                if size > MAX_BYTES:
                    raise ValueError("Input exceeds 25 MiB limit")
                digest.update(chunk)
                target.write(chunk)
        headers, rows = load_table(temp, sheet_name)
        identifier = uuid.uuid4().hex
        dest = store.path("datasets", identifier)
        dest.mkdir()
        snapshot = dest / ("source" + temp.suffix)
        os.replace(temp, snapshot)
        columns = []
        for index, header in enumerate(headers):
            numeric = missing = 0
            low, high = math.inf, -math.inf
            for row in rows:
                if not row[index].strip():
                    missing += 1
                    continue
                try:
                    value = float(row[index])
                    if math.isfinite(value):
                        numeric += 1
                        low, high = min(low, value), max(high, value)
                except ValueError:
                    pass
            columns.append(
                {
                    "name": header,
                    "numeric": numeric,
                    "missing": missing,
                    "other": len(rows) - numeric - missing,
                    "min": low if numeric else None,
                    "max": high if numeric else None,
                }
            )
        metadata = {
            "dataset_id": identifier,
            "name": source.name,
            "sha256": digest.hexdigest(),
            "snapshot": snapshot.name,
            "sheet_name": sheet_name,
            "bytes": size,
            "row_count": len(rows),
            "columns": columns,
            "preview": [row[:12] for row in rows[:4]],
            "preview_columns": headers[:12],
        }
        write_json(dest / "metadata.json", metadata)
        return metadata | {"snapshot": "local immutable snapshot"}
    finally:
        temp.unlink(missing_ok=True)


def dataset_table(store: Store, identifier: str):
    directory = store.path("datasets", identifier)
    meta = read_json(directory / "metadata.json")
    snapshot = directory / meta["snapshot"]
    if snapshot.parent != directory or sha256(snapshot) != meta["sha256"]:
        raise ValueError("Dataset snapshot integrity check failed; inspect the file again")
    return meta, *load_table(snapshot, meta["sheet_name"])


def numeric_columns(headers: list[str], rows: list[list[str]], names: list[str]) -> dict[str, list[float]]:
    result = {}
    errors = []
    for name in dict.fromkeys(names):
        if name not in headers:
            raise ValueError(f"Column {name!r} not found")
        index = headers.index(name)
        values = []
        for number, row in enumerate(rows, 2):
            try:
                value = float(row[index])
                if not math.isfinite(value):
                    raise ValueError
                values.append(value)
            except ValueError:
                if len(errors) < 8:
                    errors.append({"row": number, "column": name})
        result[name] = values
    if errors:
        raise ValueError(
            f"Selected columns contain missing/nonfinite/nonnumeric values. No rows dropped: {errors}"
        )
    return result


def column_digest(columns: list[list[float]]) -> str:
    return hashlib.sha256(json_bytes(columns)).hexdigest()
